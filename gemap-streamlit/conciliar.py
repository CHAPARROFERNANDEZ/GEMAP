#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CONCILIADOR DE MAYORES
======================
Lee el listado "Cuentas corrientes" (mayor) exportado de la contabilidad y genera un
Excel con:
  - tu mayor original + columnas L:P (Estado, Confianza, Grupo, Propuesta/comentario)
  - hoja Resumen (saldos por proveedor, con formulas)
  - hoja Incidencias (saldos deudores, posibles duplicados, cuentas repetidas...)
  - hoja Pendientes proveedores
  - hoja Propuestas 555 (donde deberia ir cada apunte de la cuenta 555)

Uso:
    python conciliar.py mayor1.xlsx [mayor2.xlsx ...]
    python conciliar.py carpeta_con_mayores
Opciones:
    --prov 40,41      prefijos de cuentas de proveedores/acreedores
    --cli 43          prefijos de cuentas de clientes
    --partidas 55500000   cuenta de partidas pendientes de aplicacion
    --salida CARPETA  carpeta donde guardar los resultados

Las propuestas son SUGERENCIAS basadas en importes, fechas y referencias.
Validalas antes de contabilizar nada.
"""
import argparse
import itertools
import os
import re
import sys
from collections import defaultdict
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "1.0"

# ----------------------------------------------------------------------------
# Utilidades
# ----------------------------------------------------------------------------

def eur(c):
    return round(c / 100, 2)


def fmt(c):
    return f"{c / 100:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


STOP = {"SU", "FRA", "FRAS", "S.L", "SL", "SA", "TRANSFER", "TRANSF", "PAGO", "LIMITED",
        "LTD", "GROUP", "SERVEIS", "DE", "LA", "EL", "S.A", "SLU", "EUROS", "INMEDIATA",
        "PAGOS", "EU", "S.L.U"}


def digs(s):
    t = str(s).replace(".", "")
    return {d.lstrip("0") for d in re.findall(r"\d{4,}", t)}


def toks(s):
    return {t for t in re.findall(r"[A-Za-z0-9]{4,}", str(s).upper()) if t not in STOP}


SUFIJOS = re.compile(r"\b(S\.?L\.?U?\.?|S\.?A\.?U?\.?|S\.?L\.?L\.?|SLP|SCP|S\.?COOP\.?|LTD\.?|LIMITED|SRL|S\.?R\.?L\.?)\b",
                     re.I)


def norm_nombre(n):
    n = str(n).upper()
    n = SUFIJOS.sub(" ", n)
    n = re.sub(r"[^A-Z0-9ÁÉÍÓÚÑ ]", " ", n)
    return re.sub(r"\s+", " ", n).strip()


# ----------------------------------------------------------------------------
# 1. Lectura del mayor
# ----------------------------------------------------------------------------

def leer_mayor(ruta):
    wb = load_workbook(ruta, read_only=True, data_only=False)
    ws = wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    empresa, fecha_listado = os.path.basename(ruta), None
    inicio = None
    for i, r in enumerate(filas[:40]):
        a = r[0]
        if isinstance(a, str):
            if a.startswith("Empresa:"):
                empresa = a.split(":", 1)[1].strip()
            m = re.match(r"Fecha:\s*(\d\d/\d\d/\d{4})", a)
            if m:
                fecha_listado = datetime.strptime(m.group(1), "%d/%m/%Y")
            if a.strip() == "Cuenta":
                inicio = i + 1
                break
    if inicio is None:
        raise ValueError("No encuentro la fila de cabecera ('Cuenta', 'Descripción'...). "
                         "¿Es un listado 'Cuentas corrientes'?")
    acc = desc = None
    recs = []
    for i in range(inicio, len(filas)):
        r = list(filas[i]) + [None] * 10
        a, b = r[0], r[1]
        es_num = isinstance(a, (int, float)) or (isinstance(a, str) and a.strip().isdigit())
        if es_num and b is not None and r[3] is None and r[4] is None:
            acc, desc = int(a), str(b)
            continue
        if r[4] is None and r[6] is None:
            continue
        if r[4] == "Total cuenta" or acc is None:
            continue
        recs.append((i + 1, acc, desc) + tuple(r[2:10]))
    df = pd.DataFrame(recs, columns=["row", "cta", "desc", "punt", "fecha", "concepto",
                                     "doc", "debe", "haber", "saldo", "contra"])
    for c in ("debe", "haber"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    df["debe_c"] = (df.debe * 100).round().astype(int)
    df["haber_c"] = (df.haber * 100).round().astype(int)
    df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce")
    df["doc"] = df["doc"].fillna("").astype(str)
    df["concepto"] = df["concepto"].fillna("").astype(str)
    return df, empresa, fecha_listado or datetime.today(), ws.title


# ----------------------------------------------------------------------------
# 2. Casado dentro de cada cuenta
# ----------------------------------------------------------------------------

def dt(x):
    return x if pd.notna(x) else pd.Timestamp("2025-12-31")


def casar_cuenta(sub, tipo):
    """Devuelve lista de (ids_documento, ids_liquidacion, 'Exacto'|'Combinado')."""
    sgn = -1 if tipo == "sup" else 1  # sup: factura = haber (net<0) ; cli: factura = debe (net>0)
    D = [(i, sgn * sub.at[i, "net"], dt(sub.at[i, "fecha"])) for i in sub.index if sgn * sub.at[i, "net"] > 0]
    S = [(i, -sgn * sub.at[i, "net"], dt(sub.at[i, "fecha"])) for i in sub.index if -sgn * sub.at[i, "net"] > 0]
    uD, uS, res = set(), set(), []
    by = defaultdict(list)
    for d in D:
        by[d[1]].append(d)
    for s in sorted(S, key=lambda x: x[2]):
        c = [d for d in by.get(s[1], []) if d[0] not in uD]
        if not c:
            continue
        c.sort(key=lambda d: abs((s[2] - d[2]).days) + (0 if d[2] <= s[2] else 1000))
        d = c[0]
        uD.add(d[0]); uS.add(s[0])
        res.append(([d[0]], [s[0]], "Exacto"))
    rD = [d for d in D if d[0] not in uD]
    rS = [s for s in S if s[0] not in uS]
    if len(rD) <= 45 and len(rS) <= 45:
        for A, B, swap in ((rS, rD, False), (rD, rS, True)):
            for a in sorted(A, key=lambda x: -x[1]):
                usedA = uD if swap else uS
                usedB = uS if swap else uD
                if a[0] in usedA:
                    continue
                av = [b for b in B if b[0] not in usedB]
                for k in (2, 3):
                    hit = None
                    for comb in itertools.combinations(av, k):
                        if sum(b[1] for b in comb) == a[1]:
                            hit = comb
                            break
                    if hit:
                        ids = [b[0] for b in hit]
                        usedA.add(a[0]); usedB.update(ids)
                        res.append(((ids, [a[0]], "Combinado") if not swap else ([a[0]], ids, "Combinado")))
                        break
    return res


def analizar(df, prov, cli, partidas):
    df = df.copy()
    df["pref"] = df.cta.astype(str)
    df["net"] = df.debe_c - df.haber_c
    for c in ("estado", "grupo", "sug", "conf"):
        df[c] = ""
    es_sup = df.pref.str[:2].isin(prov)
    es_cli = df.pref.str[:2].isin(cli)
    cnt = [0]

    def ng(p):
        cnt[0] += 1
        return f"{p}{cnt[0]:04d}"

    for tipo, mask in (("sup", es_sup), ("cli", es_cli)):
        for cta, sub in df[mask].groupby("cta"):
            for a, b, tp in casar_cuenta(sub, tipo):
                g = ng("C" if tp == "Exacto" else "K")
                for i in a + b:
                    df.at[i, "estado"] = "Casado" if tp == "Exacto" else "Casado (varios)"
                    df.at[i, "grupo"] = g
            for i in sub.index:
                if df.at[i, "estado"] == "":
                    df.at[i, "estado"] = "Pendiente"

    def short(i):
        r = df.loc[i]
        return f"{r['desc'].title()[:28]} · {str(r['concepto'])[:22]} {r['doc'] or ''}".strip()

    # --- cruces entre proveedores (mismo importe, signo contrario, distinta cuenta)
    pend = df[es_sup & (df.estado == "Pendiente")]
    used = set()
    by = defaultdict(list)
    for i, r in pend.iterrows():
        by[abs(r.net)].append(i)
    for amt, ids in by.items():
        neg = [i for i in ids if df.at[i, "net"] < 0]
        pos = [i for i in ids if df.at[i, "net"] > 0]
        for p in pos:
            c = [n for n in neg if n not in used and df.at[n, "cta"] != df.at[p, "cta"]]
            if not c or p in used:
                continue
            n = min(c, key=lambda n: abs((dt(df.at[p, "fecha"]) - dt(df.at[n, "fecha"])).days))
            used.update([p, n])
            g = ng("X")
            conf = "Alta" if amt >= 50000 or toks(df.at[p, "concepto"]) & toks(df.at[n, "desc"]) else "Media"
            for a, b in ((p, n), (n, p)):
                df.at[a, "estado"] = "Propuesto"; df.at[a, "grupo"] = g; df.at[a, "conf"] = conf
                df.at[a, "sug"] = f"Cruce entre proveedores: {fmt(amt)} € casa con {short(b)}"

    # --- cuenta 555
    m5 = df[df.cta == partidas]
    rows5 = [(i, df.at[i, "net"], dt(df.at[i, "fecha"])) for i in m5.index]
    used5 = set()
    for a in rows5:
        for b in rows5:
            if a[0] < b[0] and a[1] == -b[1] and a[0] not in used5 and b[0] not in used5:
                used5.update([a[0], b[0]])
                g = ng("P")
                for x, y in ((a[0], b[0]), (b[0], a[0])):
                    df.at[x, "estado"] = "Propuesto"; df.at[x, "grupo"] = g; df.at[x, "conf"] = "Media"
                    df.at[x, "sug"] = (f"Se compensa con otro apunte de la 555 ({fmt(abs(a[1]))} €, "
                                       f"{df.at[y, 'fecha']:%d/%m/%Y})")
    cand = df[(es_sup | es_cli) & (df.estado == "Pendiente")]
    cand_by_cta = {c: list(s.index) for c, s in cand.groupby("cta")}
    free5 = [r for r in rows5 if r[0] not in used5]
    h = defaultdict(list)
    for k in (1, 2):
        for comb in itertools.combinations(free5, k):
            h[sum(x[1] for x in comb)].append(comb)
    proposals = []
    if h:
        for cta, ids in cand_by_cta.items():
            kmax = 2 if len(ids) > 260 else 3
            for k in range(1, kmax + 1):
                if len(ids) > 60 and k > 2:
                    continue
                for comb in itertools.combinations(ids, k):
                    s = sum(df.at[i, "net"] for i in comb)
                    for c5 in h.get(-s, []):
                        proposals.append((k, len(c5), cta, comb, c5))
    proposals.sort(key=lambda p: (p[0] + p[1], p[2]))
    usedc = set()
    for k, m, cta, comb, c5 in proposals:
        if any(i in usedc for i in comb) or any(x[0] in used5 for x in c5):
            continue
        ids5 = [x[0] for x in c5]
        t5 = set().union(*[toks(df.at[i, "concepto"]) | toks(df.at[i, "doc"]) for i in ids5])
        tA = set().union(*[toks(df.at[i, "desc"]) | toks(df.at[i, "doc"]) | toks(df.at[i, "concepto"]) for i in comb])
        d5s = set().union(*[digs(df.at[i, "concepto"]) | digs(df.at[i, "doc"]) for i in ids5])
        dAs = set().union(*[digs(df.at[i, "doc"]) | digs(df.at[i, "concepto"]) for i in comb])
        dig = any(a in b or b in a for a in d5s for b in dAs if a and b)
        text = bool(t5 & tA) or dig
        nacc = len(cand_by_cta[cta])
        if (k > 1 and nacc > 15 and not text) or (m > 1 and not text):
            continue
        d5 = max(dt(df.at[i, "fecha"]) for i in ids5)
        dA = min(dt(df.at[i, "fecha"]) for i in comb)
        gap = (d5 - dA).days
        amt = sum(abs(x[1]) for x in c5)
        if dig:
            conf = "Alta"
        elif k == 1 and m == 1:
            conf = "Alta" if text or (-10 <= gap <= 120 and amt >= 10000) else "Media"
        elif m == 1 and k > 1:
            conf = "Alta" if text else "Media"
        else:
            conf = "Media" if text else "Baja"
        if amt < 2000 and not text:
            conf = "Baja"
        g = ng("P")
        for i in ids5:
            used5.add(i)
            df.at[i, "estado"] = "Propuesto"; df.at[i, "grupo"] = g; df.at[i, "conf"] = conf
            df.at[i, "sug"] = ("Imputar a " + df.at[comb[0], "desc"].title() + f" ({int(df.at[comb[0], 'cta'])})"
                               + " → cierra: " + " + ".join(f"{short(j)} {fmt(abs(df.at[j, 'net']))}" for j in comb))
        for j in comb:
            usedc.add(j)
            df.at[j, "estado"] = "Propuesto"; df.at[j, "grupo"] = g; df.at[j, "conf"] = conf
            df.at[j, "sug"] = ("Se cierra con apunte(s) de la 555: " + " + ".join(
                f"{df.at[i, 'fecha']:%d/%m/%Y} {str(df.at[i, 'concepto']).strip()} {fmt(abs(df.at[i, 'net']))}"
                for i in ids5))
    for i in rows5:
        if i[0] not in used5:
            df.at[i[0], "estado"] = "Pendiente"
    return df


# ----------------------------------------------------------------------------
# 3. Comentarios, incidencias
# ----------------------------------------------------------------------------

def tipo_pend(r, prov, cli):
    if r.estado != "Pendiente":
        return ""
    p = str(r.cta)[:2]
    if p in prov:
        if r.concepto == "Saldos anteriores":
            return "Saldo anterior sin movimiento"
        return "Factura sin pago" if r.net < 0 else "Pago/abono sin factura aplicada"
    if p in cli:
        return "Factura sin cobro" if r.net > 0 else "Cobro/abono sin factura aplicada"
    return "Sin propuesta"


def nota_555(c):
    c = str(c).upper()
    if "REMESA" in c:
        return "Posible remesa de cobros de clientes: pedir el detalle de la remesa para aplicarla a las facturas."
    if any(k in c for k in ("AMORTIZ", "PRESTAMO", "PRÉSTAMO", "TARJ")):
        return "Posible amortización de préstamo/tarjeta: revisar si va a la cuenta de deudas (52x)."
    if any(k in c for k in ("TRASPASO", "TRANSF INTERN")):
        return "Posible traspaso entre bancos: buscar el apunte contrario en las cuentas 572."
    if any(k in c for k in ("GENERALITAT", "HACIENDA", "AEAT", "TGSS", "SEGURIDAD SOCIAL", "AGENCIA TRIB")):
        return "Posible organismo público (subvención, devolución o impuesto): revisar documento de origen."
    if c.startswith(("FRA", "FACTURA", "PEDV")):
        return "Salida con referencia de factura pero sin factura pendiente del mismo importe: falta registrar la factura o el pago es de otro proveedor."
    return "Sin coincidencia por importe: revisar el concepto en el extracto bancario."


def comentarios(df, prov, cli, partidas, hoy):
    df["tipo"] = df.apply(lambda r: tipo_pend(r, prov, cli), axis=1)

    def com(r):
        if r.estado == "Casado":
            return "Casado con " + r.grupo
        if r.estado == "Casado (varios)":
            return "Casado por suma de varios apuntes (grupo " + r.grupo + ")"
        if r.estado == "Propuesto":
            return r.sug
        if r.estado == "Pendiente":
            if r.cta == partidas:
                return nota_555(r.concepto)
            if r.tipo == "Factura sin pago":
                d = (hoy - r.fecha).days if pd.notna(r.fecha) else None
                return f"Factura sin pago registrado ({d} días)" + (" - ANTIGUA >90 días" if d and d > 90 else "")
            return r.tipo
        return ""

    df["coment"] = df.apply(com, axis=1)
    return df


def incidencias(df, prov, cli, partidas, hoy):
    """Lista de (categoria, cuenta, tercero, fecha, concepto, doc, importe, comentario)."""
    out = []
    sup = df[df.pref.str[:2].isin(prov)]
    # 1 saldo deudor en proveedores
    g = sup.groupby(["cta", "desc"]).net.sum()
    for (cta, d), v in g.items():
        if v > 100:
            out.append(("Saldo DEUDOR en proveedor", cta, d, None, "", "", eur(v),
                        "Se ha pagado más de lo facturado: falta alguna factura o hay un pago mal imputado."))
    # 2 posibles duplicados: mismo importe y fecha en la misma cuenta Y (mismo documento, documento
    #   contenido en otro, o sin documento). Facturas con nº distinto se consideran distintas.
    f = sup[(sup.haber_c > 0) & (sup.concepto != "Saldos anteriores")]
    for (cta, fecha, imp), grp in f.groupby(["cta", "fecha", "haber_c"]):
        if len(grp) < 2:
            continue
        docs = [d.strip().upper() for d in grp.doc]
        sospechoso = False
        for a, b in itertools.combinations(docs, 2):
            if a == b or (a and b and (a in b or b in a)):
                sospechoso = True
        if sospechoso:
            out.append(("Posible factura duplicada", cta, grp.desc.iloc[0], fecha, "; ".join(grp.concepto.str.strip().str[:25]),
                        "; ".join(grp.doc), eur(imp) * len(grp),
                        f"{len(grp)} facturas del mismo importe y fecha con documento igual o parecido."))
    d2 = f[(f.doc.str.len() > 3)]
    for (cta, doc, imp), grp in d2.groupby(["cta", "doc", "haber_c"]):
        if len(grp) > 1 and len(set(grp.fecha)) > 1:
            out.append(("Mismo nº de factura repetido", cta, grp.desc.iloc[0], grp.fecha.iloc[0], "", doc,
                        eur(imp) * len(grp), f"El documento {doc} aparece {len(grp)} veces con el mismo importe en fechas distintas."))
    # 3 facturas a 0
    z = sup[(sup.debe_c == 0) & (sup.haber_c == 0) & (sup.concepto != "Saldos anteriores")]
    for _, r in z.iterrows():
        out.append(("Apunte de importe 0", r.cta, r.desc, r.fecha, r.concepto, r.doc, 0.0, "Apunte sin importe."))
    # 4 saldos anteriores sin mover
    for _, r in sup[(sup.concepto == "Saldos anteriores") & (sup.estado == "Pendiente")].iterrows():
        out.append(("Saldo anterior sin mover", r.cta, r.desc, None, "Saldos anteriores", "", eur(r.net),
                    "Arrastrado de ejercicios anteriores y sin ningún movimiento en el año."))
    # 5 mismo tercero con varias cuentas
    nm = sup.groupby("cta").desc.first().to_frame()
    nm["n"] = nm.desc.map(norm_nombre)
    for n, grp in nm[nm.n.str.len() > 3].groupby("n"):
        if len(grp) > 1:
            saldo = sup[sup.cta.isin(grp.index)].net.sum()
            out.append(("Tercero con varias cuentas", ", ".join(str(c) for c in grp.index), grp.desc.iloc[0], None, "", "",
                        eur(saldo), "Mismo nombre en cuentas distintas: valorar unificar."))
    # 6 mismo tercero como proveedor y como cliente
    cl = df[df.pref.str[:2].isin(cli)]
    ncl = cl.groupby("cta").desc.first().to_frame()
    ncl["n"] = ncl.desc.map(norm_nombre)
    ncl = ncl[ncl.n.str.len() > 3]
    for n, grp in nm[nm.n.str.len() > 3].groupby("n"):
        match = ncl[ncl.n == n]
        if len(match):
            sp = sup[sup.cta.isin(grp.index)].net.sum()
            sc = cl[cl.cta.isin(match.index)].net.sum()
            if abs(sp) > 100 and abs(sc) > 100:
                out.append(("Proveedor y cliente a la vez", f"{grp.index[0]} / {match.index[0]}", grp.desc.iloc[0], None, "", "",
                            eur(sp + sc), f"Saldo proveedor {fmt(sp)} € y saldo cliente {fmt(sc)} €: posible compensación."))
    # 7 cobros en un cliente que coinciden con facturas de otro cliente
    cp = cl[cl.estado == "Pendiente"]
    pos = cp[cp.net > 0]
    neg = cp[cp.net < 0]
    usados = set()
    cuentas_genericas = cl.groupby("cta").size().sort_values(ascending=False).head(1).index.tolist()
    for i, a in neg.iterrows():
        if a.cta in cuentas_genericas:
            continue
        m = pos[(pos.net == -a.net) & (pos.cta != a.cta) & (~pos.index.isin(usados))]
        if len(m):
            b = m.iloc[0]
            usados.add(m.index[0])
            out.append(("Cobro en cliente distinto", f"{a.cta} → {b.cta}", f"{a.desc} / {b.desc}", a.fecha, a.concepto, b.doc,
                        eur(abs(a.net)), "Un cobro/abono pendiente en un cliente coincide con una factura pendiente de otro."))
    # 8 cuentas sin ningún pago
    gp = sup.groupby("cta").agg(d=("debe_c", "sum"), h=("haber_c", "sum"), n=("row", "size"), desc=("desc", "first"))
    sin = gp[(gp.d == 0) & (gp.h > 0)]
    if len(sin) >= 20:
        out.append(("Cuentas de proveedor sin ningún pago", f"{len(sin)} cuentas", "(varias)", None, "", "", eur(sin.h.sum()),
                    "Muchas cuentas con facturas y ningún pago: revisar cómo se pagan (otra cuenta, efectivo, banco sin pasar...)."))
    return out


# ----------------------------------------------------------------------------
# 4. Excel de salida
# ----------------------------------------------------------------------------

HDR = PatternFill("solid", fgColor="1F3864")
HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
A9 = Font(name="Arial", size=9)
A10 = Font(name="Arial", size=10)
B10 = Font(name="Arial", size=10, bold=True)
FILLS = {"Casado": PatternFill("solid", fgColor="C6EFCE"), "Casado (varios)": PatternFill("solid", fgColor="E2F0D9"),
         "Propuesto": PatternFill("solid", fgColor="FFEB9C"), "Pendiente": PatternFill("solid", fgColor="FFC7CE")}
PROP_MED = PatternFill("solid", fgColor="FFD966")


def cab(ws, fila, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(fila, j, v)
        c.fill = HDR; c.font = HF
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def escribir_excel(ruta_origen, ruta_salida, df, empresa, hoy, prov, cli, partidas, incs):
    wb = load_workbook(ruta_origen)
    ws = wb.worksheets[0]
    SH = "'" + ws.title.replace("'", "''") + "'"
    scope = df[(df.pref.str[:2].isin(prov)) | (df.pref.str[:2].isin(cli)) | (df.cta == partidas)]
    # cabecera de las columnas nuevas
    fila_cab = int(df.row.min()) - 1
    for r in range(1, 40):
        if ws.cell(r, 1).value == "Cuenta":
            fila_cab = r
            break
    for j, h in enumerate(["Cuenta (aux)", "Estado", "Confianza", "Grupo", "Propuesta / comentario"], start=12):
        c = ws.cell(fila_cab, j, h)
        c.fill = HDR; c.font = HF; c.alignment = Alignment(horizontal="center")
    for _, r in scope.iterrows():
        rr = int(r.row)
        ws.cell(rr, 12, int(r.cta)); ws.cell(rr, 13, r.estado)
        ws.cell(rr, 14, r.conf or None); ws.cell(rr, 15, r.grupo or None); ws.cell(rr, 16, r.coment)
        f = FILLS[r.estado]
        if r.estado == "Propuesto" and r.conf in ("Media", "Baja"):
            f = PROP_MED
        for col in (5, 7, 8, 12, 13, 14, 15, 16):
            ws.cell(rr, col).fill = f
        for col in range(12, 17):
            ws.cell(rr, col).font = A9
    for col, w in zip("LMNOP", (13, 16, 11, 9, 110)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(fila_cab + 1, 1)
    ws.auto_filter.ref = f"L{fila_cab}:P{ws.max_row}"

    # ---- Resumen
    rs = wb.create_sheet("Resumen", 0)
    rs["A1"] = f"{empresa} - Conciliación de proveedores y cuenta 555 (mayor a {hoy:%d/%m/%Y})"
    rs["A1"].font = Font(name="Arial", size=13, bold=True)
    rs["A2"] = ("Estados en la hoja original (columnas L:P): Casado = factura y pago de igual importe; Casado (varios) = suma de "
                "varios apuntes; Propuesto = propuesta a validar; Pendiente = sin contrapartida.")
    rs["A2"].font = A10
    cab(rs, 4, ["Estado", "Proveedores (nº apuntes)", "Clientes (nº apuntes)", "Cuenta 555 (nº apuntes)"])
    lo_p = int(min(prov) + "000000"); hi_p = int(str(int(max(prov)) + 1) + "000000")
    lo_c = int(min(cli) + "000000"); hi_c = int(str(int(max(cli)) + 1) + "000000")
    for k, e in enumerate(["Casado", "Casado (varios)", "Propuesto", "Pendiente"], start=5):
        rs.cell(k, 1, e).font = A10; rs.cell(k, 1).fill = FILLS[e]
        rs.cell(k, 2, f'=COUNTIFS({SH}!$M:$M,$A{k},{SH}!$L:$L,">={lo_p}",{SH}!$L:$L,"<{hi_p}")')
        rs.cell(k, 3, f'=COUNTIFS({SH}!$M:$M,$A{k},{SH}!$L:$L,">={lo_c}",{SH}!$L:$L,"<{hi_c}")')
        rs.cell(k, 4, f'=COUNTIFS({SH}!$M:$M,$A{k},{SH}!$L:$L,{partidas})')
        for c in (2, 3, 4):
            rs.cell(k, c).font = A10
    rs.cell(9, 1, "Total").font = B10
    for c in (2, 3, 4):
        rs.cell(9, c, f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}8)").font = B10

    sup = scope[scope.pref.str[:2].isin(prov)]
    cab(rs, 12, ["Cuenta", "Proveedor", "Debe (pagos)", "Haber (facturas)", "Saldo (debe-haber)",
                 "Nº apuntes pendientes", "Importe pendiente (debe-haber)", "Nº propuestos", "Alerta"])
    rs.row_dimensions[12].height = 32
    pend = sup[sup.estado == "Pendiente"]
    names = sup.groupby("cta")["desc"].first()
    order = sorted(sup.cta.unique())
    for k, cta in enumerate(order):
        r = 13 + k
        g = sup[sup.cta == cta]
        pn = pend[pend.cta == cta]
        a = []
        if g.net.sum() > 100:
            a.append("Saldo DEUDOR (se ha pagado más de lo facturado)")
        fs = pn[pn.tipo == "Factura sin pago"]
        if len(fs) and (fs.fecha < hoy - pd.Timedelta(days=90)).any():
            a.append("Facturas sin pago >90 días")
        if (pn.tipo == "Saldo anterior sin movimiento").any():
            a.append("Saldo anterior sin mover")
        if (g.estado == "Propuesto").any():
            a.append("Tiene cruces/propuestas")
        rs.cell(r, 1, int(cta)); rs.cell(r, 2, names[cta])
        rs.cell(r, 3, f"=SUMIFS({SH}!$G:$G,{SH}!$L:$L,$A{r})")
        rs.cell(r, 4, f"=SUMIFS({SH}!$H:$H,{SH}!$L:$L,$A{r})")
        rs.cell(r, 5, f"=C{r}-D{r}")
        rs.cell(r, 6, f'=COUNTIFS({SH}!$L:$L,$A{r},{SH}!$M:$M,"Pendiente")')
        rs.cell(r, 7, f'=SUMIFS({SH}!$G:$G,{SH}!$L:$L,$A{r},{SH}!$M:$M,"Pendiente")-SUMIFS({SH}!$H:$H,{SH}!$L:$L,$A{r},{SH}!$M:$M,"Pendiente")')
        rs.cell(r, 8, f'=COUNTIFS({SH}!$L:$L,$A{r},{SH}!$M:$M,"Propuesto")')
        rs.cell(r, 9, "; ".join(a))
        for c in range(1, 10):
            rs.cell(r, c).font = A10
        for c in (3, 4, 5, 7):
            rs.cell(r, c).number_format = "#,##0.00;[Red]-#,##0.00"
    last = 12 + len(order)
    rs.cell(last + 1, 2, "TOTAL").font = B10
    for c in (3, 4, 5, 6, 7, 8):
        L = get_column_letter(c)
        rs.cell(last + 1, c, f"=SUM({L}13:{L}{last})").font = B10
        if c in (3, 4, 5, 7):
            rs.cell(last + 1, c).number_format = "#,##0.00;[Red]-#,##0.00"
    for col, w in zip("ABCDEFGHI", (16, 38, 16, 16, 18, 14, 20, 12, 70)):
        rs.column_dimensions[col].width = w
    rs.freeze_panes = "A13"

    # ---- Incidencias
    ic = wb.create_sheet("Incidencias", 1)
    cab(ic, 1, ["Tipo", "Cuenta", "Tercero", "Fecha", "Concepto", "Documento", "Importe", "Comentario"])
    for k, t in enumerate(incs, start=2):
        for j, v in enumerate(t, 1):
            if j == 4 and v is not None and not isinstance(v, str):
                v = pd.Timestamp(v).to_pydatetime() if pd.notna(v) else None
            ic.cell(k, j, v).font = A9
        ic.cell(k, 4).number_format = "dd/mm/yyyy"
        ic.cell(k, 7).number_format = "#,##0.00;[Red]-#,##0.00"
    for col, w in zip("ABCDEFGH", (34, 22, 36, 11, 30, 18, 14, 90)):
        ic.column_dimensions[col].width = w
    ic.freeze_panes = "A2"
    ic.auto_filter.ref = f"A1:H{max(len(incs) + 1, 2)}"

    # ---- Pendientes proveedores
    pp = wb.create_sheet("Pendientes proveedores")
    cab(pp, 1, ["Fila en hoja original", "Cuenta", "Proveedor", "Fecha", "Concepto", "Documento",
                "Importe (debe + / haber -)", "Tipo", "Antigüedad (días)", "Estado", "Confianza", "Propuesta / comentario"])
    rows = sup[sup.estado.isin(["Pendiente", "Propuesto"])].sort_values(["cta", "fecha"])
    for k, (_, r) in enumerate(rows.iterrows(), start=2):
        vals = [int(r.row), int(r.cta), r.desc, r.fecha.to_pydatetime() if pd.notna(r.fecha) else None, r.concepto,
                r.doc or None, eur(r.net), r.tipo or "Con propuesta",
                (hoy - r.fecha).days if pd.notna(r.fecha) else None, r.estado, r.conf or None, r.coment]
        for j, v in enumerate(vals, 1):
            pp.cell(k, j, v).font = A9
        pp.cell(k, 4).number_format = "dd/mm/yyyy"
        pp.cell(k, 7).number_format = "#,##0.00;[Red]-#,##0.00"
        pp.cell(k, 10).fill = FILLS[r.estado]
    for col, w in zip("ABCDEFGHIJKL", (9, 11, 34, 11, 30, 16, 16, 28, 11, 12, 10, 100)):
        pp.column_dimensions[col].width = w
    pp.freeze_panes = "A2"
    pp.auto_filter.ref = f"A1:L{max(len(rows) + 1, 2)}"

    # ---- Propuestas 555
    p5 = wb.create_sheet("Propuestas 555")
    cab(p5, 1, ["Fila en hoja original", "Fecha", "Concepto", "Documento", "Contrapartida banco", "Debe", "Haber",
                "Estado", "Confianza", "Dónde debería ir / comentario"])
    s5 = scope[scope.cta == partidas].sort_values("fecha")
    for k, (_, r) in enumerate(s5.iterrows(), start=2):
        vals = [int(r.row), r.fecha.to_pydatetime() if pd.notna(r.fecha) else None, str(r.concepto).strip(), r.doc or None,
                int(r.contra) if pd.notna(r.contra) else None, r.debe or None, r.haber or None, r.estado, r.conf or None, r.coment]
        for j, v in enumerate(vals, 1):
            p5.cell(k, j, v).font = A9
        p5.cell(k, 2).number_format = "dd/mm/yyyy"
        for c in (6, 7):
            p5.cell(k, c).number_format = "#,##0.00"
        f = FILLS[r.estado]
        if r.estado == "Propuesto" and r.conf != "Alta":
            f = PROP_MED
        p5.cell(k, 8).fill = f
    n5 = len(s5) + 1
    p5.cell(n5 + 1, 3, "Totales").font = B10
    p5.cell(n5 + 1, 6, f"=SUM(F2:F{n5})").font = B10
    p5.cell(n5 + 1, 7, f"=SUM(G2:G{n5})").font = B10
    p5.cell(n5 + 2, 3, "Saldo 555 (debe-haber)").font = B10
    p5.cell(n5 + 2, 6, f"=F{n5 + 1}-G{n5 + 1}").font = B10
    p5.cell(n5 + 3, 3, "Saldo con propuestas Alta").font = B10
    p5.cell(n5 + 3, 6, f'=SUMIFS(F2:F{n5},I2:I{n5},"Alta")-SUMIFS(G2:G{n5},I2:I{n5},"Alta")').font = B10
    for c in (6, 7):
        for rr in (n5 + 1, n5 + 2, n5 + 3):
            p5.cell(rr, c).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJ", (9, 11, 26, 18, 13, 12, 12, 12, 10, 130)):
        p5.column_dimensions[col].width = w
    p5.freeze_panes = "A2"
    wb.save(ruta_salida)


# ----------------------------------------------------------------------------
# 5. Resumen de todas las empresas
# ----------------------------------------------------------------------------

def actualizar_resumen_global(carpeta, fila):
    ruta = os.path.join(carpeta, "RESUMEN_EMPRESAS.xlsx")
    cols = ["Empresa", "Archivo", "Fecha del mayor", "Apuntes proveedores", "Casados", "Con propuesta", "Pendientes",
            "Importe pendiente proveedores", "Apuntes 555", "Saldo 555", "Nº incidencias", "Procesado el"]
    filas = []
    if os.path.exists(ruta):
        try:
            w = load_workbook(ruta)
            for r in w.active.iter_rows(min_row=2, values_only=True):
                if r[1] and r[1] != fila[1]:
                    filas.append(list(r))
        except Exception:
            pass
    filas.append(fila)
    filas.sort(key=lambda x: str(x[0]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    cab(ws, 1, cols)
    for k, r in enumerate(filas, start=2):
        for j, v in enumerate(r, 1):
            ws.cell(k, j, v).font = A10
        for c in (8, 10):
            ws.cell(k, c).number_format = "#,##0.00;[Red]-#,##0.00"
    for col, w in zip("ABCDEFGHIJKL", (38, 40, 14, 12, 10, 12, 11, 18, 11, 14, 12, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(ruta)
    return ruta


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------

def procesar(ruta, prov, cli, partidas, carpeta_salida):
    t0 = datetime.now()
    df, empresa, hoy, hoja = leer_mayor(ruta)
    hoy = pd.Timestamp(hoy)
    df = analizar(df, prov, cli, partidas)
    df = comentarios(df, prov, cli, partidas, hoy)
    incs = incidencias(df, prov, cli, partidas, hoy)
    base = os.path.splitext(os.path.basename(ruta))[0]
    salida = os.path.join(carpeta_salida, f"{base}_conciliacion.xlsx")
    escribir_excel(ruta, salida, df, empresa, hoy, prov, cli, partidas, incs)
    sup = df[df.pref.str[:2].isin(prov)]
    p5 = df[df.cta == partidas]
    ec = sup.estado.value_counts()
    fila = [empresa, os.path.basename(ruta), hoy.to_pydatetime().strftime("%d/%m/%Y"), len(sup),
            int(ec.get("Casado", 0) + ec.get("Casado (varios)", 0)), int(ec.get("Propuesto", 0)), int(ec.get("Pendiente", 0)),
            eur(sup[sup.estado == "Pendiente"].net.sum()), len(p5), eur(p5.net.sum()), len(incs),
            datetime.now().strftime("%d/%m/%Y %H:%M")]
    actualizar_resumen_global(carpeta_salida, fila)
    print(f"\n=== {empresa} ({os.path.basename(ruta)}) - {(datetime.now() - t0).seconds}s")
    print(f"  Proveedores: {len(sup)} apuntes | casados {fila[4]} | con propuesta {fila[5]} | pendientes {fila[6]}")
    print(f"  Importe pendiente proveedores: {fmt(round(fila[7] * 100))} €")
    print(f"  Cuenta 555: {len(p5)} apuntes | saldo {fmt(round(fila[9] * 100))} €")
    print(f"  Incidencias detectadas: {len(incs)}")
    print(f"  Resultado: {salida}")
    return salida


def main():
    ap = argparse.ArgumentParser(description="Conciliador de mayores (proveedores, clientes y cuenta 555)")
    ap.add_argument("entradas", nargs="+", help="archivos .xlsx del mayor o carpetas que los contengan")
    ap.add_argument("--prov", default="40,41")
    ap.add_argument("--cli", default="43")
    ap.add_argument("--partidas", type=int, default=55500000)
    ap.add_argument("--salida", default=None)
    a = ap.parse_args()
    prov = [x.strip() for x in a.prov.split(",") if x.strip()]
    cli = [x.strip() for x in a.cli.split(",") if x.strip()]
    archivos = []
    for e in a.entradas:
        if os.path.isdir(e):
            for f in sorted(os.listdir(e)):
                if f.lower().endswith(".xlsx") and "_conciliacion" not in f and not f.startswith(("~", "RESUMEN_EMPRESAS")):
                    archivos.append(os.path.join(e, f))
        else:
            archivos.append(e)
    if not archivos:
        print("No he encontrado ningún archivo .xlsx para procesar.")
        sys.exit(1)
    errores = 0
    for f in archivos:
        carpeta = a.salida or os.path.join(os.path.dirname(os.path.abspath(f)), "resultados")
        os.makedirs(carpeta, exist_ok=True)
        try:
            procesar(f, prov, cli, a.partidas, carpeta)
        except Exception as ex:  # no parar si falla una empresa
            errores += 1
            print(f"\n*** ERROR con {f}: {ex}")
    print(f"\nTerminado. Archivos procesados: {len(archivos) - errores} | con error: {errores}")


if __name__ == "__main__":
    main()
